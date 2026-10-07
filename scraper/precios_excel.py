"""
Excel de precios de Alexandra.

  python scraper/precios_excel.py exportar   # crea/actualiza "Precios Alemakeup.xlsx" (conserva lo ya llenado)
  python scraper/precios_excel.py importar   # lee la columna "Precio Alexandra" → precios_alexandra.json

El Excel tiene los costos del mayorista: es PRIVADO y no se sube a internet.
precios_alexandra.json solo tiene precios de venta (que ya son públicos en la página) y sí se publica.
"""

import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
EXCEL = RAIZ / "Precios Alemakeup.xlsx"
REPORTE = RAIZ / "scraper" / "reporte_precios.csv"
JSON_PRECIOS = RAIZ / "precios_alexandra.json"

HOJA = "Precios"
COLUMNAS = [  # (título, ancho)
    ("ID", 10),
    ("Categoría", 20),
    ("Subcategoría", 26),
    ("Producto", 60),
    ("Precio mayorista", 15),
    ("Precio detal del mayorista", 17),
    ("Precio con la regla (50 %)", 17),
    ("Precio Alexandra", 17),
    ("Precio final en la página", 17),
    ("Ganancia $", 14),
    ("Ganancia %", 12),
]
COL_ID, COL_ALEXANDRA = 1, 8

FUENTE = "Arial"
ROSA = "B4706F"
AMARILLO = PatternFill("solid", fgColor="FFF2B3")
GRIS = PatternFill("solid", fgColor="F3ECE9")
BORDE = Border(bottom=Side(style="thin", color="E2CFC9"))
PESOS = '"$"#,##0;[Red]-"$"#,##0;"-"'


def leer_reporte():
    if not REPORTE.exists():
        sys.exit("No existe scraper/reporte_precios.csv. Primero ejecuta: python scraper/sync.py")
    with open(REPORTE, encoding="utf-8-sig") as f:
        filas = list(csv.DictReader(f, delimiter=";"))
    return sorted(filas, key=lambda r: (r["grupo"], r["categoria"], r["producto"].lower()))


def precios_llenados():
    """Lo que Alexandra ya escribió: primero del Excel, si no existe, del JSON publicado."""
    precios = {}
    if JSON_PRECIOS.exists():
        precios.update({k: v["precio"] for k, v in json.loads(JSON_PRECIOS.read_text(encoding="utf-8")).items()})
    if EXCEL.exists():
        ws = load_workbook(EXCEL)[HOJA]
        for fila in ws.iter_rows(min_row=2, values_only=True):
            pid, valor = fila[COL_ID - 1], fila[COL_ALEXANDRA - 1]
            if pid:
                precios.pop(pid, None)
                if isinstance(valor, (int, float)) and valor > 0:
                    precios[pid] = int(round(valor))
    return precios


def exportar():
    filas = leer_reporte()
    previos = precios_llenados()

    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    titulo = Font(name=FUENTE, bold=True, color="FFFFFF")
    for c, (nombre, ancho) in enumerate(COLUMNAS, 1):
        celda = ws.cell(row=1, column=c, value=nombre)
        celda.font = titulo
        celda.fill = PatternFill("solid", fgColor="FFB000" if c == COL_ALEXANDRA else ROSA)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = ancho
    ws.row_dimensions[1].height = 36
    ws.cell(row=1, column=COL_ALEXANDRA).font = Font(name=FUENTE, bold=True, color="000000")
    ws.cell(row=1, column=COL_ALEXANDRA).comment = Comment(
        "Escribe aquí el precio al público que quieres cobrar (solo el número, sin $ ni puntos). "
        "Déjalo vacío para usar el precio con la regla del 50 %.", "Alemakeup")

    normal = Font(name=FUENTE, size=10)
    entrada = Font(name=FUENTE, size=10, bold=True, color="0000FF")
    for i, r in enumerate(filas, 2):
        valores = [r["id"], r["grupo"], r["categoria"], r["producto"],
                   int(r["costo_mayorista"]), int(r["precio_detal_mayorista"]), int(r["precio_regla"]),
                   previos.get(r["id"])]
        for c, v in enumerate(valores, 1):
            ws.cell(row=i, column=c, value=v).font = normal
        ws.cell(row=i, column=9, value=f'=IF(ISNUMBER(H{i}),H{i},G{i})')
        ws.cell(row=i, column=10, value=f"=I{i}-E{i}")
        ws.cell(row=i, column=11, value=f"=IF(E{i}>0,J{i}/E{i},0)")
        for c in range(1, len(COLUMNAS) + 1):
            celda = ws.cell(row=i, column=c)
            celda.border = BORDE
            celda.font = entrada if c == COL_ALEXANDRA else normal
            if 5 <= c <= 10:
                celda.number_format = PESOS
        ws.cell(row=i, column=11).number_format = '0%;[Red]-0%;"-"'
        ws.cell(row=i, column=COL_ALEXANDRA).fill = AMARILLO
        ws.cell(row=i, column=9).fill = GRIS

    ultima = len(filas) + 1
    # Rojo si el precio de Alexandra no cubre el costo
    ws.conditional_formatting.add(
        f"H2:H{ultima}",
        FormulaRule(formula=["AND(ISNUMBER(H2),H2<=E2)"],
                    fill=PatternFill("solid", start_color="F4B6B6", end_color="F4B6B6"),
                    font=Font(color="9C0006", bold=True)))
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{ultima}"
    ws.column_dimensions["A"].hidden = True

    ins = wb.create_sheet("Instrucciones", 0)
    ins.column_dimensions["A"].width = 110
    textos = [
        ("Cómo poner los precios", Font(name=FUENTE, bold=True, size=14, color=ROSA)),
        ("", None),
        ("1. Ve a la hoja «Precios».", None),
        ("2. En la columna amarilla «Precio Alexandra» escribe el precio al público que quieres cobrar.", None),
        ("   Escribe solo el número: 25000 (sin $, sin puntos).", None),
        ("3. Si dejas la casilla vacía, el producto se vende con la regla del 50 % (columna «Precio con la regla»).", None),
        ("4. Guarda el Excel y ciérralo. Luego doble clic en «Subir precios.bat».", None),
        ("", None),
        ("Ejemplo de una fila llenada:", Font(name=FUENTE, bold=True)),
        ("   Producto: Pestañinas prosa 151  ·  Precio mayorista: $11.500  ·  Precio Alexandra: 23000  →  Ganancia: $11.500 (100 %)", None),
        ("", None),
        ("Qué significa cada columna", Font(name=FUENTE, bold=True)),
        ("• Precio mayorista: lo que le cuesta el producto al mayorista (Geal Makeup).", None),
        ("• Precio detal del mayorista: lo que el mayorista les cobra a las clientas finales (referencia de mercado).", None),
        ("• Precio con la regla: el precio automático (50 % de ganancia; Milagros a su precio público).", None),
        ("• Precio final en la página: el que verán las clientas (el tuyo si lo pusiste, si no el de la regla).", None),
        ("• Ganancia $ y %: se calculan solas.", None),
        ("", None),
        ("Importante", Font(name=FUENTE, bold=True)),
        ("• Si una casilla se pone ROJA, el precio que escribiste es igual o menor al costo: estarías perdiendo.", None),
        ("• Si el mayorista sube un costo por encima de tu precio, la página usa la regla del 50 % para no vender a pérdida.", None),
        ("• Puedes usar los filtros de la fila de títulos para ver solo una categoría.", None),
        ("• Este archivo tiene los costos del mayorista: NO lo compartas ni lo publiques.", None),
        ("• Cuando lleguen productos nuevos, «Excel de precios.bat» los agrega sin borrar lo que ya llenaste.", None),
    ]
    for i, (t, f) in enumerate(textos, 1):
        celda = ins.cell(row=i, column=1, value=t)
        celda.font = f or Font(name=FUENTE, size=11)
    wb.active = 1

    try:
        wb.save(EXCEL)
    except PermissionError:
        sys.exit(f"No se pudo guardar: cierra «{EXCEL.name}» en Excel y vuelve a intentar.")
    print(f"Excel listo: {EXCEL}  ({len(filas)} productos, {len(previos)} con precio de Alexandra)")


def importar():
    if not EXCEL.exists():
        sys.exit(f"No existe «{EXCEL.name}». Primero ejecuta «Excel de precios.bat».")
    try:
        ws = load_workbook(EXCEL)[HOJA]
    except PermissionError:
        sys.exit(f"Cierra «{EXCEL.name}» en Excel y vuelve a intentar.")
    precios, perdida, raros = {}, [], []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        pid, producto, costo, valor = fila[0], fila[3], fila[4], fila[COL_ALEXANDRA - 1]
        if not pid or valor in (None, ""):
            continue
        if isinstance(valor, str):
            limpio = valor.replace("$", "").replace(".", "").replace(",", "").strip()
            if not limpio.isdigit():
                raros.append(f"{producto}: «{valor}»")
                continue
            valor = int(limpio)
        valor = int(round(valor))
        if valor <= 0:
            continue
        if costo and valor <= costo:
            perdida.append(f"{producto}: ${valor:,} (costo ${costo:,})")
        precios[pid] = {"precio": valor, "producto": producto}
    JSON_PRECIOS.write_text(json.dumps(dict(sorted(precios.items(), key=lambda x: x[1]["producto"].lower())),
                                       ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"{len(precios)} precios de Alexandra guardados en {JSON_PRECIOS.name}")
    if perdida:
        print(f"\nAVISO: {len(perdida)} precios son iguales o menores al costo (se usará la regla del 50 %):")
        print("\n".join("  - " + x for x in perdida))
    if raros:
        print(f"\nAVISO: {len(raros)} casillas no son un número y se ignoraron:")
        print("\n".join("  - " + x for x in raros))


if __name__ == "__main__":
    accion = sys.argv[1] if len(sys.argv) > 1 else ""
    if accion == "exportar":
        exportar()
    elif accion == "importar":
        importar()
    else:
        sys.exit(__doc__)
